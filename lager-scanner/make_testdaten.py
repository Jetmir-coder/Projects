"""Erzeugt Testdaten (JSON) + Test-Excel mit denselben Daten."""
import json, random
from openpyxl import load_workbook

random.seed(7)
NAMEN = ["Schraube M8x40", "Mutter M8", "Unterlegscheibe 8,4", "Kabelbinder 200mm", "Folie 500mm",
         "Karton 600x400", "Klebeband braun", "Etikett 100x50", "Dichtung DN50", "Rohr 1/2 Zoll",
         "Winkel 90 Grad", "Scharnier 60mm", "Griff Edelstahl", "Kleber 2K", "Palettenfolie",
         "Steckdose 3-fach", "Kabel 3x1,5", "Schlauch 10mm", "Filter G4", "Lager 6204"]
mats = []
for i in range(50):
    tail = random.choice(["0530", "0012", "0777"]) if i % 7 == 0 else f"{random.randint(1000, 9999)}"
    sap = f"40{random.randint(10, 99)}{tail}"
    stk_k = random.choice([0, 6, 12, 24, 50, 100])
    stk_p = (stk_k or 1) * random.choice([20, 30, 40, 60])
    mats.append({"sap": sap, "bez": f"{NAMEN[i % 20]} Typ {chr(65 + i // 20)}",
                 "stand": stk_p * random.randint(1, 4), "melde": stk_p,
                 "stkKarton": stk_k, "stkPalette": stk_p})
# 200 Stellplätze im Plan (10 Reihen x 20), davon ~70 belegt
plaetze = {p: "" for p in range(1, 201)}
for p in random.sample(range(1, 201), 70):
    plaetze[p] = random.choice(mats)["sap"]
json.dump({"mats": mats, "plaetze": plaetze}, open("test/testdaten.json", "w"), ensure_ascii=False, indent=1)

wb = load_workbook("Lager_Vorlage.xlsx")
def fill(ws, tname, rows):
    ws.delete_rows(2, ws.max_row)
    for r in rows:
        ws.append(r)
    t = ws.tables[tname]
    t.ref = f"A1:{t.ref.split(':')[1].rstrip('0123456789')}{1 + len(rows)}"
fill(wb["Materialliste"], "tblMaterial",
     [[m["sap"], m["bez"], m["stand"], m["melde"], m["stkKarton"], m["stkPalette"]] for m in mats])
fill(wb["Stellplaetze"], "tblStellplatz", [[p, s] for p, s in plaetze.items()])
lay = wb["PlanLayout"]
for r in range(2, 42):
    for c in range(2, 42):
        lay.cell(r, c).value = None
# 10 Reihen à 20 Plätze, nach jeder 2. Reihe ein Gang
row = 3
for reihe in range(10):
    for c in range(20):
        lay.cell(row, 2 + c, reihe * 20 + c + 1)
    row += 1 if reihe % 2 == 0 else 2
wb.save("test/Lager_Test.xlsx")
print(len(mats), "Materialien,", sum(1 for s in plaetze.values() if s), "belegt")
