from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font

wb = Workbook()

def tabelle(ws, name, header, rows=()):
    ws.append(header)
    for r in rows:
        ws.append(r)
    last = max(ws.max_row, 2)
    t = Table(displayName=name, ref=f"A1:{chr(64+len(header))}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for i, h in enumerate(header):
        ws.column_dimensions[chr(65+i)].width = max(14, len(h) + 4)

ws = wb.active
ws.title = "Materialliste"
tabelle(ws, "tblMaterial",
        ["SAP", "Bezeichnung", "Lagerstand", "Meldebestand", "StkProKarton", "StkProPalette"],
        [["1000001", "Beispielmaterial", 0, 10, 12, 480]])
for r in ws.iter_rows(min_row=2, max_col=1):
    r[0].number_format = "@"

ws2 = wb.create_sheet("Stellplaetze")
tabelle(ws2, "tblStellplatz", ["Stellplatz", "SAP"], [[1, ""]])
for r in ws2.iter_rows(min_row=2, max_col=2):
    r[1].number_format = "@"

ws3 = wb.create_sheet("Protokoll")
tabelle(ws3, "tblProtokoll",
        ["Zeit", "Stellplatz", "SAP", "Aktion", "Einheit", "Anzahl", "Stueck", "NeuerBestand"],
        [["", "", "", "", "", "", "", ""]])

ws4 = wb.create_sheet("Lageruebersicht")
ws4.append(["Stellplatz", "SAP", "Bezeichnung", "Lagerstand (alle Plaetze)", "Meldebestand", "Unter Meldebestand?"])
for c in ws4[1]:
    c.font = Font(bold=True)
for i in range(2, 1002):
    ws4.cell(i, 1, f"=IFERROR(INDEX(tblStellplatz[Stellplatz],ROW()-1),\"\")")
    ws4.cell(i, 2, f"=IFERROR(INDEX(tblStellplatz[SAP],ROW()-1),\"\")")
    ws4.cell(i, 3, f"=IF(B{i}=\"\",\"\",IFERROR(INDEX(tblMaterial[Bezeichnung],MATCH(B{i},tblMaterial[SAP],0)),\"?\"))")
    ws4.cell(i, 4, f"=IF(B{i}=\"\",\"\",IFERROR(INDEX(tblMaterial[Lagerstand],MATCH(B{i},tblMaterial[SAP],0)),\"?\"))")
    ws4.cell(i, 5, f"=IF(B{i}=\"\",\"\",IFERROR(INDEX(tblMaterial[Meldebestand],MATCH(B{i},tblMaterial[SAP],0)),\"?\"))")
    ws4.cell(i, 6, f"=IF(OR(B{i}=\"\",NOT(ISNUMBER(D{i}))),\"\",IF(D{i}<E{i},\"JA\",\"\"))")
for col, w in zip("ABCDEF", (12, 14, 30, 24, 14, 20)):
    ws4.column_dimensions[col].width = w

wb.save("Lager_Vorlage.xlsx")
print("ok")
